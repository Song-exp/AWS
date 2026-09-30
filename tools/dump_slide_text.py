# -*- coding: utf-8 -*-
"""지정한 pptx의 특정 슬라이드 전체 텍스트를 위치순으로 덤프 (읽기 전용).

사용: python tools/dump_slide_text.py <pptx경로> [슬라이드번호 ...]
"""

import sys
from pptx import Presentation

EMU = 914400


def dump(path, wanted):
    prs = Presentation(path)
    for idx, slide in enumerate(prs.slides, 1):
        if wanted and idx not in wanted:
            continue
        print("=" * 76)
        print("슬라이드 %d  (도형 %d개)" % (idx, len(slide.shapes)))
        print("=" * 76)
        rows = []
        for sh in slide.shapes:
            if sh.has_table:
                for r, row in enumerate(sh.table.rows):
                    for c, cl in enumerate(row.cells):
                        t = cl.text_frame.text.strip()
                        if t:
                            rows.append((sh.top or 0, sh.left or 0,
                                         "[표 %d행%d열] %s" % (r, c, t.replace("\n", " / "))))
                continue
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if t:
                rows.append((sh.top or 0, sh.left or 0, t.replace("\n", " / ")))
        rows.sort(key=lambda x: (round(x[0] / EMU, 2), round(x[1] / EMU, 2)))
        for top, left, t in rows:
            print("  y=%.2f x=%.2f  %s" % (top / EMU, left / EMU, t))
        print()


if __name__ == "__main__":
    p = sys.argv[1]
    nums = {int(a) for a in sys.argv[2:]} if len(sys.argv) > 2 else set()
    dump(p, nums)
