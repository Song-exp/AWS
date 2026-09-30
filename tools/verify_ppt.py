# -*- coding: utf-8 -*-
"""생성된 pptx 구조 검증. python tools/verify_ppt.py"""

import os
import sys

from pptx import Presentation
from pptx.util import Emu

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def newest_deck():
    """가장 최근에 저장된 pptx (잠긴 구버전이 옆에 남아 있을 수 있음)."""
    decks = [os.path.join(ROOT, f) for f in os.listdir(ROOT) if f.lower().endswith(".pptx")]
    if not decks:
        raise SystemExit("pptx 파일이 없습니다: %s" % ROOT)
    return max(decks, key=os.path.getmtime)


DECK = newest_deck()

EMU_IN = 914400


def main():
    prs = Presentation(DECK)
    sw, sh = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    print(f"파일       : {os.path.basename(DECK)}")
    print(f"슬라이드 크기 : {sw:.3f} x {sh:.3f} in")
    print(f"슬라이드 수  : {len(prs.slides)}")
    print()

    problems = []
    for idx, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        tables = [s for s in shapes if s.has_table]
        n_txt = sum(1 for s in shapes if s.has_text_frame and s.text_frame.text.strip())
        n_runs = 0
        fonts = set()
        for s in shapes:
            if s.has_text_frame:
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        n_runs += 1
                        fonts.add(r.font.name)
        cells_txt = 0
        for gf in tables:
            for row in gf.table.rows:
                for c in row.cells:
                    if c.text_frame.text.strip():
                        cells_txt += 1
                        for p in c.text_frame.paragraphs:
                            for r in p.runs:
                                fonts.add(r.font.name)

        # 슬라이드 경계 이탈 검사
        for s in shapes:
            if s.left is None or s.top is None:
                continue
            l, t = s.left / EMU_IN, s.top / EMU_IN
            r = l + (s.width or 0) / EMU_IN
            b = t + (s.height or 0) / EMU_IN
            if l < -0.02 or t < -0.02 or r > sw + 0.02 or b > sh + 0.02:
                problems.append(
                    f"  슬라이드 {idx}: 경계 이탈 {s.shape_type} "
                    f"L{l:.2f} T{t:.2f} R{r:.2f} B{b:.2f}")

        grid = [f"{len(gf.table.rows)}x{len(gf.table.columns)}" for gf in tables]
        print(f"슬라이드 {idx}: 도형 {len(shapes):3d} | 텍스트도형 {n_txt:3d} | "
              f"런 {n_runs:3d} | 표 {len(tables)}개 {grid} | 표셀텍스트 {cells_txt}")
        print(f"           폰트: {sorted(f for f in fonts if f)}")

    print()
    if problems:
        print("[문제 발견]")
        print("\n".join(problems))
        return 1
    print("[OK] 모든 도형이 슬라이드 경계 안에 있음")
    return 0


if __name__ == "__main__":
    sys.exit(main())
