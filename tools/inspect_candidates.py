# -*- coding: utf-8 -*-
"""복구 후보 pptx들의 내용을 읽어 비교한다 (읽기 전용)."""

import os
from pptx import Presentation

HOME = os.path.expanduser("~")
KAKAO = os.path.join(HOME, "Documents", "카카오톡 받은 파일")
WS = r"C:\Users\bella\Desktop\대학\공모전\AWS"

CANDIDATES = [
    ("A. 카톡 11:16 수정본", os.path.join(KAKAO, "PayPick_발표자료_수정본.pptx")),
    ("B. 카톡 12:40 최신", os.path.join(KAKAO, "PayPick_발표자료.pptx")),
    ("C. NetShortcuts 10:56", os.path.join(
        HOME, "AppData", "Roaming", "Microsoft", "Windows",
        "Network Shortcuts", "PayPick_발표자료.pptx")),
    ("D. 현재 TMI (사용자 12:31 편집)", os.path.join(WS, "TMI_발표자료.pptx")),
]


def head_text(slide, limit=3):
    """슬라이드에서 위쪽 텍스트 몇 줄을 뽑아 식별용으로 보여준다."""
    items = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip().replace("\n", " / ")
        if t:
            items.append((sh.top if sh.top is not None else 0, t))
    items.sort(key=lambda x: x[0])
    return [t for _, t in items[:limit]]


for label, path in CANDIDATES:
    print("=" * 78)
    if not os.path.exists(path):
        print(label, "-> 파일 없음")
        continue
    size = os.path.getsize(path)
    print("%s  |  %s bytes" % (label, format(size, ",")))
    print(path)
    try:
        prs = Presentation(path)
    except Exception as e:
        print("  열기 실패:", e)
        continue
    print("  슬라이드 %d장 / %.2f x %.2f in"
          % (len(prs.slides), prs.slide_width / 914400, prs.slide_height / 914400))
    for i, s in enumerate(prs.slides, 1):
        n_tbl = sum(1 for sh in s.shapes if sh.has_table)
        n_pic = sum(1 for sh in s.shapes if sh.shape_type == 13)
        lines = head_text(s)
        print("   [%d] 도형%3d 표%d 이미지%d" % (i, len(s.shapes), n_tbl, n_pic))
        for ln in lines:
            print("        ", ln[:88])
